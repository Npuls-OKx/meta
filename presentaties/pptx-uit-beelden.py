#!/usr/bin/env python3
"""Bouwt een PowerPoint uit de geëxporteerde slidebeelden, met de sprekersnotities erbij.

Elke slide is het gerenderde beeld op volle grootte; de tekst is daardoor niet
bewerkbaar, maar de slide ziet er overal precies zo uit als in de browser. De
route via LibreOffice (pdf naar pptx) levert wel bewerkbare tekst, maar vervormt
kaarten met een gekleurde rand en eigen figuren: de kaart krijgt de randkleur als
vulling en een tekening kan de hele slide zwart maken. Voor een deck dat gedeeld
wordt is dit script daarom de betrouwbare route.

Gebruik:
    python3 pptx-uit-beelden.py src/<deck>.md export/<deck>/img export/<deck>/<deck>.pptx
"""

import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

COMMENTAAR = re.compile(r"<!--(.*?)-->", re.DOTALL)


def notities(markdown):
    """Per slide de sprekersnotitie: het laatste commentaarblok, zoals Slidev dat doet."""
    tekst = markdown.read_text(encoding="utf-8")
    if tekst.startswith("---"):
        tekst = tekst.split("\n---\n", 1)[1]
    slides = re.split(r"\n---\n", tekst)
    uit = []
    for s in slides:
        blokken = COMMENTAAR.findall(s)
        uit.append(blokken[-1].strip() if blokken else "")
    return uit


def bouw(markdown, beeldmap, doel):
    beelden = sorted(Path(beeldmap).glob("*.png"), key=lambda p: int(p.stem))
    if not beelden:
        sys.exit(f"geen beelden gevonden in {beeldmap}; draai eerst ./deck <naam> beelden")
    noten = notities(Path(markdown))
    if len(noten) != len(beelden):
        print(f"waarschuwing: {len(noten)} slides in de markdown, {len(beelden)} beelden; "
              f"notities worden op volgorde gekoppeld", file=sys.stderr)

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    leeg = prs.slide_layouts[6]
    for i, beeld in enumerate(beelden):
        slide = prs.slides.add_slide(leeg)
        slide.shapes.add_picture(str(beeld), 0, 0, width=prs.slide_width, height=prs.slide_height)
        if i < len(noten) and noten[i]:
            slide.notes_slide.notes_text_frame.text = noten[i]
    Path(doel).parent.mkdir(parents=True, exist_ok=True)
    prs.save(doel)
    print(f"PowerPoint: {doel} ({len(beelden)} slides, {sum(1 for n in noten if n)} met notitie)")


if __name__ == "__main__":
    if len(sys.argv) != 4:
        sys.exit(__doc__)
    bouw(*sys.argv[1:])
